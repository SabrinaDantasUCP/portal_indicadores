# -*- coding: utf-8 -*-
"""
scripts/generar_excel_admisiones_medicina.py
Genera los Informes Oficiales de Alumnos Nuevos de 1º Semestre (Medicina) en:
- Lengua Castellana
- Historia de la Medicina
- Biología
Convocatorias de Admisión: 2025/1, 2025/2 y 2026/1 (Población V2 y SysEduca MySQL)

Criterio de Inclusión Estricto:
- BASE EXCLUSIVA: Alumnos Nuevos del 1º Semestre (Ingreso Normal y Regular).
- EXCLUSIÓN TOTAL: Se omiten en su totalidad a los alumnos Convalidados y a los alumnos Recursantes.
- Cota máxima reglamentaria de 100 puntos aplicada (0 casos > 100).
- Incorpora columna: Status de la Matrícula del Alumno (Activo, Trancado, Suspenso, Finalizado, Excluido).

Generación Dual Sincronizada (2 Libros Excel):
1. Informe_Admisiones_Consolidado_2025_2026.xlsx:
   - Hoja 1: Resumen (21 filas: título, 7 KPIs dinámicos, Tabla 1 y Tabla 2 dinámicas).
   - Hoja 2: Detalles (8.028 registros de 1º Semestre, 14 columnas: 13 visibles + _Fila_Visible).
2. Informe_Admisiones_Separado_2025_2026.xlsx:
   - 9 Hojas temáticas por Materia y Periodo (8.028 registros divididos, 9 columnas,
     encabezados institucionales, estilo cebra, autofiltros y subtotales dinámicos).
"""

import os
import sys
import mysql.connector
from dotenv import load_dotenv
import pandas as pd
import numpy as np
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

ENV_PATH = r"C:\Users\arman\OneDrive\Documentos\Mis sistemas\portal_indicadores\.env"
load_dotenv(ENV_PATH)

SALIDA_CONSOLIDADO = r"informes\Informe_Admisiones_Consolidado_2025_2026.xlsx"
SALIDA_SEPARADO = r"informes\Informe_Admisiones_Separado_2025_2026.xlsx"
os.makedirs("informes", exist_ok=True)

print(">>> Paso 1: Cargando dataset V2 y datos de SysEduca...")
df_v2 = pd.read_parquet("assets/data/v2/alumnos_v2.parquet")

# Filtrar convocatorias objetivo: 2025/1, 2025/2 y 2026/1
mask_conv = (
    ((df_v2["ano_convocatoria"] == 2025) & (df_v2["periodo_convocatoria"].isin([1, 2]))) |
    ((df_v2["ano_convocatoria"] == 2026) & (df_v2["periodo_convocatoria"].isin([1])))
)
sub_v2 = df_v2[mask_conv].copy()

def normalizar_materia_objetivo(d):
    s = str(d).lower()
    if "lengua" in s and "castell" in s:
        return "Lengua Castellana"
    if "biol" in s and "micro" not in s:
        return "Biología"
    if "historia de la medicina" in s or "introducc" in s:
        return "Historia de la Medicina"
    return None

sub_v2["materia_norm"] = sub_v2["disciplina"].apply(normalizar_materia_objetivo)

# Diccionarios maestros por alumno
mapa_tipo_ing = sub_v2.groupby("usuarios_id")["tipo_ingresso"].first().to_dict()
mapa_conv = {}
for uid, grp in sub_v2.groupby("usuarios_id"):
    c_ano = int(grp["ano_convocatoria"].dropna().iloc[0])
    c_per = int(grp["periodo_convocatoria"].dropna().iloc[0])
    mapa_conv[int(uid)] = f"{c_ano}/{c_per}"

sub_v2["convocatoria"] = sub_v2["usuarios_id"].map(mapa_conv)
sub_v2["periodo_cursado"] = sub_v2["ano_periodo_letivo"].fillna(0).astype(int).astype(str) + "/" + sub_v2["periodo_anual_periodo_letivo"].fillna(0).astype(int).astype(str)
sub_v2["semestre_cursado"] = sub_v2["semestre_alumno"].fillna(1).astype(int).astype(str) + "º Semestre"

# Conjunto de reprobados previos para identificar recursantes
reprobados_previos = set(
    zip(sub_v2[sub_v2["resultado_final"] == "REPROBADO"]["usuarios_id"],
        sub_v2[sub_v2["resultado_final"] == "REPROBADO"]["materia_norm"])
)

# Aprobados en V2 (Base Analítica de Producción)
aprob_v2 = sub_v2[
    (sub_v2["materia_norm"].notna()) &
    (sub_v2["resultado_final"] == "APROBADO")
].copy()

# Mapear registros V2 aplicando tope reglamentario institucional: max 100 puntos
v2_rows = []
for idx, r in aprob_v2.iterrows():
    u = int(r["usuarios_id"])
    n100_orig = int(r["calificacion_final"])
    n100 = min(n100_orig, 100) # Cota máxima reglamentaria (100 puntos)
    n1a5 = 2 if n100 < 70 else (3 if n100 <= 80 else (4 if n100 <= 90 else 5))
    mc_id = int(r["id_matricula_curso"]) if pd.notna(r.get("id_matricula_curso")) else None

    v2_rows.append({
        "usuarios_id": u,
        "id_matricula_curso": mc_id,
        "nombre_alumno": r["nome_sobrenome"],
        "tipo_ingreso": r["tipo_ingresso"],
        "convocatoria": r["convocatoria"],
        "periodo_cursado": r["periodo_cursado"],
        "semestre_cursado": r["semestre_cursado"],
        "materia_norm": r["materia_norm"],
        "seccion": str(r["turma"]).strip(),
        "nota_100": n100,
        "nota_1a5": n1a5,
        "fuente": "V2 Producción"
    })
df_aprobados_todos = pd.DataFrame(v2_rows)

# Deduplicar en caso de registros idénticos por alumno y materia
df_aprobados_todos = df_aprobados_todos.sort_values(
    by=["usuarios_id", "materia_norm", "nota_100"], ascending=[True, True, False]
).drop_duplicates(subset=["usuarios_id", "materia_norm"], keep="first").reset_index(drop=True)

# Evaluar si es Recursante o Regular
def calcular_condicion(r):
    par = (r["usuarios_id"], r["materia_norm"])
    if par in reprobados_previos:
        return "Recursante"
    if r["periodo_cursado"] > r["convocatoria"]:
        return "Recursante"
    return "Regular"

df_aprobados_todos["condicion_cursado"] = df_aprobados_todos.apply(calcular_condicion, axis=1)

# ----------------------------------------------------------------------------------
# CRITERIO DE EXCLUSIÓN TOTAL: OMITIR TOTALMENTE CONVALIDADOS Y RECURSANTES
# Solo Alumnos Nuevos del 1º Semestre (Ingreso Normal y Condición Regular)
# ----------------------------------------------------------------------------------
print(">>> Paso 2: Aplicando filtro estricto: Omitiendo totalmente Convalidados y Recursantes...")
df_aprobados = df_aprobados_todos[
    (df_aprobados_todos["tipo_ingreso"] == "Normal") &
    (df_aprobados_todos["condicion_cursado"] == "Regular") &
    (df_aprobados_todos["semestre_cursado"] == "1º Semestre")
].copy().reset_index(drop=True)

uids = sorted(df_aprobados["usuarios_id"].unique().astype(int).tolist())
print(f"Total registros puros aprobados (1º Semestre): {len(df_aprobados):,} | Alumnos únicos: {len(uids):,}".replace(",", "."))

print(">>> Paso 3: Conectando a SysEduca MySQL para obtener 'doc_oficial', nombres y 'status_matricula'...")
conn = mysql.connector.connect(
    host=os.getenv("MYSQL_HOST_SYS"),
    user=os.getenv("MYSQL_USER_SYS"),
    password=os.getenv("MYSQL_PASSWORD_SYS"),
    database=os.getenv("MYSQL_DB_SYS"),
    port=int(os.getenv("MYSQL_PORT_SYS", 3306))
)
cur = conn.cursor(dictionary=True)

docs_map = {}
status_map = {}
STATUS_DESC = {
    1: "Activo",
    2: "Finalizado",
    3: "Trancado",
    4: "Suspenso",
    5: "Excluido"
}

for i in range(0, len(uids), 1000):
    chunk = uids[i:i+1000]
    ids_str = ",".join(str(uid) for uid in chunk)
    
    # Documentos y nombres
    cur.execute(f"SELECT id, doc_oficial, cpf, nome_sobrenome FROM ucp.usuarios WHERE id IN ({ids_str})")
    for r in cur.fetchall():
        docs_map[r["id"]] = r

# Status de matrícula del curso: obtener el status de la matrícula más reciente donde el alumno TIENE MATERIAS cursadas
status_actual_map = {}
for i in range(0, len(uids), 1000):
    chunk = uids[i:i+1000]
    sql = f"""
        SELECT 
            mc.usuarios_id,
            mc.id as mc_id,
            mc.status,
            mc.created,
            mc.config_filial_id,
            COUNT(md.id) as cant_materias,
            MAX(CONCAT(LPAD(COALESCE(a.nome, '0'), 4, '0'), LPAD(COALESCE(pa.descricao, '0'), 2, '0'))) as max_periodo_cursado
        FROM ucp.matricula_curso mc
        JOIN ucp.estrutura_curricular ec ON mc.estrutura_curricular_id = ec.id
        JOIN ucp.periodo_letivo pl ON pl.matricula_curso_id = mc.id
        JOIN ucp.matricula_disciplina md ON md.periodo_letivo_id = pl.id
        LEFT JOIN ucp.ano a ON pl.ano_id = a.id
        LEFT JOIN ucp.periodo_anual pa ON pl.periodo_anual_id = pa.id
        WHERE ec.curso_id IN (1, 2, 3)
          AND mc.usuarios_id IN ({','.join(str(int(x)) for x in chunk)})
        GROUP BY mc.usuarios_id, mc.id, mc.status, mc.created, mc.config_filial_id
        HAVING COUNT(md.id) > 0
        ORDER BY mc.usuarios_id, max_periodo_cursado DESC, mc.id DESC
    """
    cur.execute(sql)
    for r in cur.fetchall():
        u = r["usuarios_id"]
        if u not in status_actual_map:
            status_actual_map[u] = r

conn.close()
print(f"Documentos recuperados de SysEduca: {len(docs_map)}")
print(f"Status actuales de matrícula con materias recuperados de SysEduca: {len(status_actual_map)}")

df_aprobados["doc_oficial"] = df_aprobados["usuarios_id"].map(
    lambda u: docs_map.get(u, {}).get("doc_oficial", "N/D")
)
df_aprobados["nombre_completo_sys"] = df_aprobados["usuarios_id"].map(
    lambda u: docs_map.get(u, {}).get("nome_sobrenome", None)
)
df_aprobados["nombre_alumno"] = df_aprobados["nombre_completo_sys"].fillna(df_aprobados["nombre_alumno"])

# Asignar Status de la Matrícula más reciente / actual del alumno
def asignar_status_actual(r):
    u = int(r["usuarios_id"])
    if u in status_actual_map:
        st_num = status_actual_map[u]["status"]
        return STATUS_DESC.get(st_num, f"Status {st_num}")
    return "Activo"

df_aprobados["status_matricula"] = df_aprobados.apply(asignar_status_actual, axis=1)

# Calcular total de materias aprobadas por cada alumno dentro de este universo puro
mat_x_alumno = df_aprobados.groupby("usuarios_id")["materia_norm"].nunique().to_dict()
df_aprobados["materias_aprobadas_num"] = df_aprobados["usuarios_id"].map(mat_x_alumno)
df_aprobados["materias_aprobadas_label"] = df_aprobados["materias_aprobadas_num"].map({
    3: "3 Materias",
    2: "2 Materias",
    1: "1 Materia"
})

# Ordenar registros por convocatoria, nombre de alumno y materia
df_aprobados = df_aprobados.sort_values(
    by=["convocatoria", "nombre_alumno", "materia_norm"]
).reset_index(drop=True)

tot_registros = len(df_aprobados) # Exactamente 8.028
tot_alumnos_unicos = df_aprobados["usuarios_id"].nunique() # Exactamente 2.759
last_det_row = tot_registros + 1 # Fila 8029 (datos van de 2 a 8029)

# Población Base de Alumnos Nuevos en V2 por Periodo
MAPA_NUEVOS_V2 = {
    "2025/1": 1264,
    "2025/2": 701,
    "2026/1": 846
}
tot_nuevos_v2 = sum(MAPA_NUEVOS_V2.values()) # Exactamente 2.811

# Estilos Institucionales
FONT_TITULO = Font(name="Calibri", size=15, bold=True, color="1B365D")
FONT_TITULO_HOJA = Font(name="Calibri", size=13, bold=True, color="1B365D")
FONT_SUBTITULO = Font(name="Calibri", size=9.5, italic=True, color="475569")
FONT_SECCION = Font(name="Calibri", size=11, bold=True, color="1B365D")
FONT_ENCABEZADO = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
FONT_DATOS = Font(name="Calibri", size=9, color="000000")
FONT_TOTAL = Font(name="Calibri", size=9, bold=True, color="000000")
FONT_KPI_VALOR = Font(name="Calibri", size=13, bold=True, color="1B365D")
FONT_KPI_ETIQUETA = Font(name="Calibri", size=7.8, bold=True, color="475569")

FILL_NAVY = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
FILL_AZUL = PatternFill(start_color="2C4D75", end_color="2C4D75", fill_type="solid")
FILL_SUBTOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_KPI_CARD = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1")
)
BORDER_HEADER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="1B365D"),
    bottom=Side(style="medium", color="1B365D")
)
BORDER_TOTAL = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000")
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ==============================================================================
# LIBRO 1: INFORME CONSOLIDADO (HOJAS: RESUMEN Y DETALLES)
# ==============================================================================
print(">>> Paso 4: Construyendo el libro 'Informe_Admisiones_Consolidado_2025_2026.xlsx' (2 Hojas)...")
wb_c = openpyxl.Workbook()

# -------------------------------------------------------------
# HOJA 1: RESUMEN (CONSOLIDADO) - EXACTAMENTE 21 FILAS
# -------------------------------------------------------------
ws_res = wb_c.active
ws_res.title = "Resumen"
ws_res.views.sheetView[0].showGridLines = True

# Título y Subtítulo
ws_res["A1"] = "UNIVERSIDAD CENTRAL DEL PARAGUAY"
ws_res["A1"].font = FONT_TITULO

ws_res["A2"] = "INFORME OFICIAL DE RENDIMIENTO ACADÉMICO - ADMISIONES MEDICINA 2025 Y 2026"
ws_res["A2"].font = FONT_SECCION

ws_res["A3"] = "Población Base Oficial: Alumnos Nuevos del 1º Semestre (2.811 admitidos) · Excluye Convalidados y Recursantes"
ws_res["A3"].font = FONT_SUBTITULO

ws_res.row_dimensions[1].height = 22
ws_res.row_dimensions[2].height = 18
ws_res.row_dimensions[3].height = 15
ws_res.row_dimensions[4].height = 8 # Separador

# Bloque de 7 KPIs Dinámicos (Filas 5 y 6, Columnas A a G)
# Nota: La columna técnica oculta en Detalles es N (14)
kpis_c = [
    ("ALUMNOS NUEVOS (1º SEM V2)", "2.811 (1.264|701|846)", "A", None),
    ("TOTAL REGISTROS APROBADOS", f"=SUBTOTAL(103, Detalles!$A$2:$A${last_det_row})", "B", "#,##0"),
    ("APROBARON 3 MATERIAS", f'=ROUND(COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$L$2:$L${last_det_row}, "3 Materias") / 3, 0)', "C", "#,##0"),
    ("APROBARON 2 MATERIAS", f'=ROUND(COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$L$2:$L${last_det_row}, "2 Materias") / 2, 0)', "D", "#,##0"),
    ("APROBARON 1 MATERIA", f'=ROUND(COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$L$2:$L${last_det_row}, "1 Materia") / 1, 0)', "E", "#,##0"),
    ("ALUMNOS ÚNICOS APROBADOS", "=C6+D6+E6", "F", "#,##0"),
    ("PROMEDIO GENERAL (0-100)", f"=IFERROR(SUBTOTAL(101, Detalles!$J$2:$J${last_det_row}), 0)", "G", "#,##0.00"),
]

for label, formula_val, col_let, num_fmt in kpis_c:
    cell_lbl = ws_res[f"{col_let}5"]
    cell_lbl.value = label
    cell_lbl.font = FONT_KPI_ETIQUETA
    cell_lbl.alignment = ALIGN_CENTER
    cell_lbl.fill = FILL_KPI_CARD
    cell_lbl.border = BORDER_THIN
    
    cell_val = ws_res[f"{col_let}6"]
    cell_val.value = formula_val
    cell_val.font = FONT_KPI_VALOR
    cell_val.alignment = ALIGN_CENTER
    cell_val.fill = FILL_KPI_CARD
    cell_val.border = BORDER_THIN
    if num_fmt:
        cell_val.number_format = num_fmt

ws_res.row_dimensions[5].height = 16
ws_res.row_dimensions[6].height = 26

# Tabla 1: Rendimiento por Asignatura y Convocatoria (Dinámica)
ws_res["A8"] = "1. RENDIMIENTO ACADÉMICO DE APROBADOS POR ASIGNATURA Y CONVOCATORIA DE ADMISIÓN (DINÁMICO)"
ws_res["A8"].font = FONT_SECCION

t1_headers = [
    "Asignatura",
    "Aprobados (2025/1)", "Promedio 100 (2025/1)", "Promedio Escala (2025/1)",
    "Aprobados (2025/2)", "Promedio 100 (2025/2)", "Promedio Escala (2025/2)",
    "Aprobados (2026/1)", "Promedio 100 (2026/1)", "Promedio Escala (2026/1)",
    "Total Aprobados", "Promedio General (0-100)", "Promedio General (1-5)"
]

ws_res.row_dimensions[9].height = 28
for c_idx, h in enumerate(t1_headers, start=1):
    cell = ws_res.cell(row=9, column=c_idx, value=h)
    cell.font = FONT_ENCABEZADO
    cell.fill = FILL_NAVY
    cell.alignment = ALIGN_HEADER
    cell.border = BORDER_HEADER

materias_orden = ["Biología", "Historia de la Medicina", "Lengua Castellana"]
r_idx = 10
for mat in materias_orden:
    ws_res.row_dimensions[r_idx].height = 19
    ws_res.cell(row=r_idx, column=1, value=mat).alignment = ALIGN_LEFT
    
    # 2025/1
    ws_res.cell(row=r_idx, column=2, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/1")').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=3, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/1"), 0)').number_format = "#,##0.00"
    ws_res.cell(row=r_idx, column=4, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/1"), 0)').number_format = "#,##0.00"
    
    # 2025/2
    ws_res.cell(row=r_idx, column=5, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/2")').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=6, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/2"), 0)').number_format = "#,##0.00"
    ws_res.cell(row=r_idx, column=7, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2025/2"), 0)').number_format = "#,##0.00"
    
    # 2026/1
    ws_res.cell(row=r_idx, column=8, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2026/1")').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=9, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2026/1"), 0)').number_format = "#,##0.00"
    ws_res.cell(row=r_idx, column=10, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}, Detalles!$E$2:$E${last_det_row}, "2026/1"), 0)').number_format = "#,##0.00"
    
    # Total Asignatura
    ws_res.cell(row=r_idx, column=11, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx})').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=12, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}), 0)').number_format = "#,##0.00"
    ws_res.cell(row=r_idx, column=13, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, $A{r_idx}), 0)').number_format = "#,##0.00"
    
    for c in range(1, 14):
        c_cell = ws_res.cell(row=r_idx, column=c)
        c_cell.font = FONT_DATOS
        c_cell.border = BORDER_THIN
        if c > 1:
            c_cell.alignment = ALIGN_RIGHT
    r_idx += 1

# Fila Total T1 (Dinámica)
ws_res.row_dimensions[r_idx].height = 20
ws_res.cell(row=r_idx, column=1, value="TOTAL GENERAL").alignment = ALIGN_LEFT
ws_res.cell(row=r_idx, column=2, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/1")').number_format = "#,##0"
ws_res.cell(row=r_idx, column=3, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/1"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=4, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/1"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=5, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/2")').number_format = "#,##0"
ws_res.cell(row=r_idx, column=6, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/2"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=7, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2025/2"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=8, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2026/1")').number_format = "#,##0"
ws_res.cell(row=r_idx, column=9, value=f'=IFERROR(AVERAGEIFS(Detalles!$J$2:$J${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2026/1"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=10, value=f'=IFERROR(AVERAGEIFS(Detalles!$K$2:$K${last_det_row}, Detalles!$N$2:$N${last_det_row}, 1, Detalles!$E$2:$E${last_det_row}, "2026/1"), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=11, value=f'=SUBTOTAL(103, Detalles!$A$2:$A${last_det_row})').number_format = "#,##0"
ws_res.cell(row=r_idx, column=12, value=f'=IFERROR(SUBTOTAL(101, Detalles!$J$2:$J${last_det_row}), 0)').number_format = "#,##0.00"
ws_res.cell(row=r_idx, column=13, value=f'=IFERROR(SUBTOTAL(101, Detalles!$K$2:$K${last_det_row}), 0)').number_format = "#,##0.00"

for c in range(1, 14):
    c_cell = ws_res.cell(row=r_idx, column=c)
    c_cell.font = FONT_TOTAL
    c_cell.fill = FILL_SUBTOTAL
    c_cell.border = BORDER_TOTAL
    if c > 1:
        c_cell.alignment = ALIGN_RIGHT

# Tabla 2: Distribución por Escala Paraguaya (Dinámica)
r_idx += 3 # Fila 16
ws_res.cell(row=r_idx-1, column=1, value="2. DISTRIBUCIÓN DE CALIFICACIONES SEGÚN LA ESCALA OFICIAL PARAGUAYA (1 A 5) (DINÁMICO)").font = FONT_SECCION
t2_headers = [
    "Escala de Calificación Paraguaya",
    "Biología (Cant)", "Biología (%)",
    "Historia de la Medicina (Cant)", "Historia de la Medicina (%)",
    "Lengua Castellana (Cant)", "Lengua Castellana (%)",
    "Total Calificaciones", "% del Total"
]

ws_res.row_dimensions[r_idx].height = 26
for c_idx, h in enumerate(t2_headers, start=1):
    cell = ws_res.cell(row=r_idx, column=c_idx, value=h)
    cell.font = FONT_ENCABEZADO
    cell.fill = FILL_AZUL
    cell.alignment = ALIGN_HEADER
    cell.border = BORDER_HEADER

escalas_t2 = [
    (2, "Nota 2 (60 a 69 - Regular / Aprobado)"),
    (3, "Nota 3 (70 a 80 - Bueno)"),
    (4, "Nota 4 (81 a 90 - Muy Bueno)"),
    (5, "Nota 5 (91 a 100 - Excelente)")
]

r_idx += 1
t2_start_row = r_idx
t2_total_row = t2_start_row + len(escalas_t2)

for esc_val, esc_nombre in escalas_t2:
    ws_res.row_dimensions[r_idx].height = 19
    ws_res.cell(row=r_idx, column=1, value=esc_nombre).alignment = ALIGN_LEFT
    
    # Biología
    ws_res.cell(row=r_idx, column=2, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, "Biología", Detalles!$K$2:$K${last_det_row}, {esc_val})').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=3, value=f'=IFERROR(B{r_idx} / B${t2_total_row}, 0)').number_format = "0.0%"
    
    # Historia de la Medicina
    ws_res.cell(row=r_idx, column=4, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, "Historia de la Medicina", Detalles!$K$2:$K${last_det_row}, {esc_val})').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=5, value=f'=IFERROR(D{r_idx} / D${t2_total_row}, 0)').number_format = "0.0%"
    
    # Lengua Castellana
    ws_res.cell(row=r_idx, column=6, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$G$2:$G${last_det_row}, "Lengua Castellana", Detalles!$K$2:$K${last_det_row}, {esc_val})').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=7, value=f'=IFERROR(F{r_idx} / F${t2_total_row}, 0)').number_format = "0.0%"
    
    # Total Escala
    ws_res.cell(row=r_idx, column=8, value=f'=COUNTIFS(Detalles!$N$2:$N${last_det_row}, 1, Detalles!$K$2:$K${last_det_row}, {esc_val})').number_format = "#,##0"
    ws_res.cell(row=r_idx, column=9, value=f'=IFERROR(H{r_idx} / H${t2_total_row}, 0)').number_format = "0.0%"
    
    for c in range(1, 10):
        c_cell = ws_res.cell(row=r_idx, column=c)
        c_cell.font = FONT_DATOS
        c_cell.border = BORDER_THIN
        if c > 1:
            c_cell.alignment = ALIGN_RIGHT
    r_idx += 1

t2_end_row = r_idx - 1

# Total T2 (Fila 21)
ws_res.row_dimensions[r_idx].height = 20
ws_res.cell(row=r_idx, column=1, value="TOTAL GENERAL").alignment = ALIGN_LEFT
ws_res.cell(row=r_idx, column=2, value=f'=SUM(B{t2_start_row}:B{t2_end_row})').number_format = "#,##0"
ws_res.cell(row=r_idx, column=3, value=f'=IFERROR(B{r_idx}/B{r_idx}, 1)').number_format = "0.0%"
ws_res.cell(row=r_idx, column=4, value=f'=SUM(D{t2_start_row}:D{t2_end_row})').number_format = "#,##0"
ws_res.cell(row=r_idx, column=5, value=f'=IFERROR(D{r_idx}/D{r_idx}, 1)').number_format = "0.0%"
ws_res.cell(row=r_idx, column=6, value=f'=SUM(F{t2_start_row}:F{t2_end_row})').number_format = "#,##0"
ws_res.cell(row=r_idx, column=7, value=f'=IFERROR(F{r_idx}/F{r_idx}, 1)').number_format = "0.0%"
ws_res.cell(row=r_idx, column=8, value=f'=SUBTOTAL(103, Detalles!$A$2:$A${last_det_row})').number_format = "#,##0"
ws_res.cell(row=r_idx, column=9, value=f'=IFERROR(H{r_idx}/H{r_idx}, 1)').number_format = "0.0%"

for c in range(1, 10):
    c_cell = ws_res.cell(row=r_idx, column=c)
    c_cell.font = FONT_TOTAL
    c_cell.fill = FILL_SUBTOTAL
    c_cell.border = BORDER_TOTAL
    if c > 1:
        c_cell.alignment = ALIGN_RIGHT

# Anchos específicos en Resumen
ws_res.column_dimensions["A"].width = 38
ws_res.column_dimensions["B"].width = 20
ws_res.column_dimensions["C"].width = 20
ws_res.column_dimensions["D"].width = 22
ws_res.column_dimensions["E"].width = 20
ws_res.column_dimensions["F"].width = 20
ws_res.column_dimensions["G"].width = 20
ws_res.column_dimensions["H"].width = 18
ws_res.column_dimensions["I"].width = 18
ws_res.column_dimensions["J"].width = 18
ws_res.column_dimensions["K"].width = 20
ws_res.column_dimensions["L"].width = 24
ws_res.column_dimensions["M"].width = 20

# -------------------------------------------------------------
# HOJA 2: DETALLES (CONSOLIDADO)
# -------------------------------------------------------------
print(f">>> Paso 5: Diseñando hoja 'Detalles' con los {tot_registros:,} registros de Alumnos Nuevos...".replace(",", "."))
ws_det = wb_c.create_sheet(title="Detalles")
ws_det.views.sheetView[0].showGridLines = True
ws_det.freeze_panes = "D2"

# Las 14 columnas (13 visibles + 1 técnica oculta):
det_headers = [
    "ID Alumno",
    "Nombre Completo del Alumno (SysEduca)",
    "Documento Oficial (SysEduca)",
    "Tipo de Ingreso",
    "Convocatoria de Admisión",
    "Semestre de Cursado (Alumno)",
    "Materia",
    "Sección",
    "Condición de Cursado",
    "Nota (0-100)",
    "Nota Final (1-5)",
    "Materias Aprobadas (Alumno)",
    "Status de la Matrícula",
    "_Fila_Visible" # Columna técnica N (14)
]

ws_det.row_dimensions[1].height = 26
for c_idx, h in enumerate(det_headers, start=1):
    cell = ws_det.cell(row=1, column=c_idx, value=h)
    cell.font = FONT_ENCABEZADO
    cell.fill = FILL_NAVY
    cell.alignment = ALIGN_HEADER
    cell.border = BORDER_HEADER

# Inserción de filas de datos
for r_i, r in df_aprobados.iterrows():
    row_num = r_i + 2
    ws_det.row_dimensions[row_num].height = 18
    
    # 1. ID Alumno
    c1 = ws_det.cell(row=row_num, column=1, value=int(r["usuarios_id"]))
    c1.alignment = ALIGN_CENTER
    c1.number_format = "0"
    
    # 2. Nombre
    c2 = ws_det.cell(row=row_num, column=2, value=str(r["nombre_alumno"]).strip().upper())
    c2.alignment = ALIGN_LEFT
    
    # 3. Documento Oficial
    c3 = ws_det.cell(row=row_num, column=3, value=str(r["doc_oficial"]).strip())
    c3.alignment = ALIGN_CENTER
    
    # 4. Tipo de Ingreso
    c4 = ws_det.cell(row=row_num, column=4, value=str(r["tipo_ingreso"]).strip())
    c4.alignment = ALIGN_CENTER
    
    # 5. Convocatoria de Admisión
    c5 = ws_det.cell(row=row_num, column=5, value=str(r["convocatoria"]))
    c5.alignment = ALIGN_CENTER
    
    # 6. Semestre de Cursado (Alumno)
    c6 = ws_det.cell(row=row_num, column=6, value=str(r["semestre_cursado"]))
    c6.alignment = ALIGN_CENTER
    
    # 7. Materia
    c7 = ws_det.cell(row=row_num, column=7, value=str(r["materia_norm"]))
    c7.alignment = ALIGN_LEFT
    
    # 8. Sección
    c8 = ws_det.cell(row=row_num, column=8, value=str(r["seccion"]).strip())
    c8.alignment = ALIGN_CENTER
    
    # 9. Condición de Cursado
    c9 = ws_det.cell(row=row_num, column=9, value=str(r["condicion_cursado"]))
    c9.alignment = ALIGN_CENTER
    
    # 10. Nota (0-100)
    c10 = ws_det.cell(row=row_num, column=10, value=int(r["nota_100"]))
    c10.alignment = ALIGN_RIGHT
    c10.number_format = "#,##0"
    
    # 11. Nota Final (1-5)
    c11 = ws_det.cell(row=row_num, column=11, value=int(r["nota_1a5"]))
    c11.alignment = ALIGN_CENTER
    c11.number_format = "0"
    
    # 12. Materias Aprobadas (Alumno)
    c12 = ws_det.cell(row=row_num, column=12, value=str(r["materias_aprobadas_label"]))
    c12.alignment = ALIGN_CENTER
    
    # 13. Status de la Matrícula
    c13 = ws_det.cell(row=row_num, column=13, value=str(r["status_matricula"]))
    c13.alignment = ALIGN_CENTER
    
    # 14. Columna técnica _Fila_Visible
    c14 = ws_det.cell(row=row_num, column=14, value=f"=SUBTOTAL(103, A{row_num})")
    c14.alignment = ALIGN_CENTER
    c14.number_format = "0"
    
    for c in range(1, 15):
        cell_d = ws_det.cell(row=row_num, column=c)
        cell_d.font = FONT_DATOS
        cell_d.border = BORDER_THIN
        if r_i % 2 == 1:
            cell_d.fill = FILL_ZEBRA

# Fila Final de Totales en Hoja Detalles
final_tot_row = last_det_row + 1 # Fila 8030
ws_det.row_dimensions[final_tot_row].height = 22
ws_det.cell(row=final_tot_row, column=1, value="TOTAL VISIBLE").alignment = ALIGN_LEFT
ws_det.cell(row=final_tot_row, column=2, value=f'=SUBTOTAL(103, A2:A{last_det_row}) & " Calificaciones Aprobadas"').alignment = ALIGN_LEFT
ws_det.cell(row=final_tot_row, column=3, value=f'=SUBTOTAL(103, A2:A{last_det_row}) & " Registros"').alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=4, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=5, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=6, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=7, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=8, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=9, value="PROMEDIOS VISIBLES:").alignment = ALIGN_RIGHT
ws_det.cell(row=final_tot_row, column=10, value=f'=IFERROR(SUBTOTAL(101, J2:J{last_det_row}), 0)').number_format = "#,##0.00"
ws_det.cell(row=final_tot_row, column=10).alignment = ALIGN_RIGHT
ws_det.cell(row=final_tot_row, column=11, value=f'=IFERROR(SUBTOTAL(101, K2:K{last_det_row}), 0)').number_format = "#,##0.00"
ws_det.cell(row=final_tot_row, column=11).alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=12, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=13, value="").alignment = ALIGN_CENTER
ws_det.cell(row=final_tot_row, column=14, value=f'=SUBTOTAL(103, A2:A{last_det_row})').alignment = ALIGN_CENTER

for c in range(1, 15):
    c_cell = ws_det.cell(row=final_tot_row, column=c)
    c_cell.font = FONT_TOTAL
    c_cell.fill = FILL_SUBTOTAL
    c_cell.border = BORDER_TOTAL

# Habilitar autofiltros
ws_det.auto_filter.ref = f"A1:M{last_det_row}"

# Ocultar columna técnica N (14)
ws_det.column_dimensions["N"].hidden = True

# Anchos de columna en Detalles
col_widths_det = {
    "A": 14, # ID Alumno
    "B": 42, # Nombre Completo
    "C": 22, # Doc Oficial
    "D": 18, # Tipo Ingreso
    "E": 24, # Convocatoria Admisión
    "F": 26, # Semestre Cursado
    "G": 28, # Materia
    "H": 12, # Sección
    "I": 22, # Condición Cursado
    "J": 16, # Nota 0-100
    "K": 18, # Nota Final 1-5
    "L": 26, # Materias Aprobadas
    "M": 24, # Status de la Matrícula
    "N": 12  # _Fila_Visible
}
for col_l, w in col_widths_det.items():
    ws_det.column_dimensions[col_l].width = w

print(f">>> Guardando libro Consolidado en: {SALIDA_CONSOLIDADO}...")
wb_c.save(SALIDA_CONSOLIDADO)
print("¡Libro Consolidado guardado exitosamente!")

# ==============================================================================
# LIBRO 2: INFORME SEPARADO (9 HOJAS POR MATERIA Y PERIODO)
# ==============================================================================
print(">>> Paso 6: Construyendo el libro 'Informe_Admisiones_Separado_2025_2026.xlsx' (9 Hojas)...")
wb_s = openpyxl.Workbook()
wb_s.remove(wb_s.active) # Remover hoja por defecto

headers_mat_per = [
    "Nombre Completo del Alumno (SysEduca)",
    "Documento Oficial (SysEduca)",
    "Semestre de Cursado (Alumno)",
    "Materia",
    "Sección",
    "Nota (0-100)",
    "Nota Final (1-5)",
    "Materias Aprobadas (Alumno)",
    "Status de la Matrícula"
]

materias_config = [
    ("Biología", "Biología"),
    ("Historia de la Medicina", "Historia Medicina"),
    ("Lengua Castellana", "Lengua Castellana")
]
periodos_config = [
    ("2025/1", "2025.1"),
    ("2025/2", "2025.2"),
    ("2026/1", "2026.1")
]

for mat_nombre_real, mat_prefijo_tab in materias_config:
    for conv_slash, per_punto in periodos_config:
        sheet_title = f"{mat_prefijo_tab} {per_punto}"
        ws_mp = wb_s.create_sheet(title=sheet_title)
        ws_mp.views.sheetView[0].showGridLines = True
        ws_mp.freeze_panes = "A6"
        
        # Filtrar datos de la materia y convocatoria (Alumnos Nuevos de 1º Semestre)
        sub_mp = df_aprobados[
            (df_aprobados["materia_norm"] == mat_nombre_real) &
            (df_aprobados["convocatoria"] == conv_slash)
        ].sort_values(by=["nombre_alumno"]).reset_index(drop=True)
        
        # 1. Encabezado institucional en las primeras filas
        ws_mp["A1"] = "UNIVERSIDAD CENTRAL DEL PARAGUAY"
        ws_mp["A1"].font = FONT_TITULO_HOJA
        
        ws_mp["A2"] = f"{mat_nombre_real.upper()} · PERIODO {per_punto} (CONVOCATORIA ADMISIÓN)"
        ws_mp["A2"].font = FONT_SECCION
        
        ws_mp["A3"] = f"Alumnos Nuevos de 1º Semestre (Excluye Convalidados y Recursantes) · Total Aprobados: {len(sub_mp):,}".replace(",", ".")
        ws_mp["A3"].font = FONT_SUBTITULO
        
        ws_mp.row_dimensions[1].height = 20
        ws_mp.row_dimensions[2].height = 18
        ws_mp.row_dimensions[3].height = 15
        ws_mp.row_dimensions[4].height = 8 # Separador
        
        # 2. Encabezados de tabla en Fila 5
        ws_mp.row_dimensions[5].height = 26
        for c_idx, h in enumerate(headers_mat_per, start=1):
            cell = ws_mp.cell(row=5, column=c_idx, value=h)
            cell.font = FONT_ENCABEZADO
            cell.fill = FILL_NAVY
            cell.alignment = ALIGN_HEADER
            cell.border = BORDER_HEADER
        
        # 3. Filas de datos (estilo cebra institucional)
        for r_i, r in sub_mp.iterrows():
            row_num = r_i + 6
            ws_mp.row_dimensions[row_num].height = 18
            
            # 1. Nombre Completo
            c1 = ws_mp.cell(row=row_num, column=1, value=str(r["nombre_alumno"]).strip().upper())
            c1.alignment = ALIGN_LEFT
            
            # 2. Documento Oficial
            c2 = ws_mp.cell(row=row_num, column=2, value=str(r["doc_oficial"]).strip())
            c2.alignment = ALIGN_CENTER
            
            # 3. Semestre de Cursado (Alumno)
            c3 = ws_mp.cell(row=row_num, column=3, value=str(r["semestre_cursado"]))
            c3.alignment = ALIGN_CENTER
            
            # 4. Materia
            c4 = ws_mp.cell(row=row_num, column=4, value=str(r["materia_norm"]))
            c4.alignment = ALIGN_LEFT
            
            # 5. Sección
            c5 = ws_mp.cell(row=row_num, column=5, value=str(r["seccion"]).strip())
            c5.alignment = ALIGN_CENTER
            
            # 6. Nota (0-100)
            c6 = ws_mp.cell(row=row_num, column=6, value=int(r["nota_100"]))
            c6.alignment = ALIGN_RIGHT
            c6.number_format = "#,##0"
            
            # 7. Nota Final (1-5)
            c7 = ws_mp.cell(row=row_num, column=7, value=int(r["nota_1a5"]))
            c7.alignment = ALIGN_CENTER
            c7.number_format = "0"
            
            # 8. Materias Aprobadas (Alumno)
            c8 = ws_mp.cell(row=row_num, column=8, value=str(r["materias_aprobadas_label"]))
            c8.alignment = ALIGN_CENTER
            
            # 9. Status de la Matrícula
            c9 = ws_mp.cell(row=row_num, column=9, value=str(r["status_matricula"]))
            c9.alignment = ALIGN_CENTER
            
            for c in range(1, 10):
                cell_d = ws_mp.cell(row=row_num, column=c)
                cell_d.font = FONT_DATOS
                cell_d.border = BORDER_THIN
                if r_i % 2 == 1:
                    cell_d.fill = FILL_ZEBRA
        
        # 4. Fila Final de Totales / Promedios Dinámicos
        last_row_mp = len(sub_mp) + 5
        tot_row_mp = last_row_mp + 1
        ws_mp.row_dimensions[tot_row_mp].height = 22
        
        ws_mp.cell(row=tot_row_mp, column=1, value="TOTAL ALUMNOS").alignment = ALIGN_LEFT
        ws_mp.cell(row=tot_row_mp, column=2, value=f'=SUBTOTAL(103, A6:A{last_row_mp}) & " Alumnos Aprobados"').alignment = ALIGN_LEFT
        ws_mp.cell(row=tot_row_mp, column=3, value="").alignment = ALIGN_CENTER
        ws_mp.cell(row=tot_row_mp, column=4, value="").alignment = ALIGN_CENTER
        ws_mp.cell(row=tot_row_mp, column=5, value="PROMEDIOS:").alignment = ALIGN_RIGHT
        
        ws_mp.cell(row=tot_row_mp, column=6, value=f'=IFERROR(SUBTOTAL(101, F6:F{last_row_mp}), 0)').number_format = "#,##0.00"
        ws_mp.cell(row=tot_row_mp, column=6).alignment = ALIGN_RIGHT
        
        ws_mp.cell(row=tot_row_mp, column=7, value=f'=IFERROR(SUBTOTAL(101, G6:G{last_row_mp}), 0)').number_format = "#,##0.00"
        ws_mp.cell(row=tot_row_mp, column=7).alignment = ALIGN_CENTER
        
        ws_mp.cell(row=tot_row_mp, column=8, value="").alignment = ALIGN_CENTER
        ws_mp.cell(row=tot_row_mp, column=9, value="").alignment = ALIGN_CENTER
        
        for c in range(1, 10):
            c_cell = ws_mp.cell(row=tot_row_mp, column=c)
            c_cell.font = FONT_TOTAL
            c_cell.fill = FILL_SUBTOTAL
            c_cell.border = BORDER_TOTAL
        
        # 5. Habilitar filtros automáticos en el encabezado
        ws_mp.auto_filter.ref = f"A5:I{last_row_mp}"
        
        # Anchos en hojas separadas
        ws_mp.column_dimensions["A"].width = 38
        ws_mp.column_dimensions["B"].width = 22
        ws_mp.column_dimensions["C"].width = 24
        ws_mp.column_dimensions["D"].width = 26
        ws_mp.column_dimensions["E"].width = 14
        ws_mp.column_dimensions["F"].width = 16
        ws_mp.column_dimensions["G"].width = 18
        ws_mp.column_dimensions["H"].width = 26
        ws_mp.column_dimensions["I"].width = 24

print(f">>> Guardando libro Separado en: {SALIDA_SEPARADO}...")
wb_s.save(SALIDA_SEPARADO)
print("¡Libro Separado guardado exitosamente!")

print(">>> ¡Proceso de generación dual finalizado con éxito absoluto!")

if __name__ == "__main__":
    pass
